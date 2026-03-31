from __future__ import annotations

import csv
import io
import json
import mimetypes
import shutil
import subprocess
import zipfile
from email import policy
from email.parser import BytesParser
from email.message import Message, EmailMessage
from html import unescape
from mailbox import mbox
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any, Iterator

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from .utils import clean_text

try:
    import extract_msg
except Exception:  # pragma: no cover
    extract_msg = None

try:
    import pytesseract
    from PIL import Image
except Exception:  # pragma: no cover
    pytesseract = None
    Image = None


SUPPORTED_EMAIL_EXTENSIONS = {".pst", ".eml", ".msg", ".mbox"}
SUPPORTED_DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".xlsm",
    ".csv",
    ".txt",
    ".md",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".png",
    ".jpg",
    ".jpeg",
    ".tif",
    ".tiff",
    ".bmp",
    ".zip",
}


def guess_mime_type(path: Path) -> str | None:
    content_type, _ = mimetypes.guess_type(path.name)
    return content_type


def html_to_text(value: str | bytes | None) -> str:
    if value is None:
        return ""
    if isinstance(value, bytes):
        value = value.decode("utf-8", errors="ignore")
    soup = BeautifulSoup(value, "html.parser")
    text = soup.get_text("\n")
    return clean_text(unescape(text))


def extract_text_from_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_text(path)
    if suffix == ".docx":
        return extract_docx_text(path)
    if suffix in {".xlsx", ".xlsm"}:
        return extract_xlsx_text(path)
    if suffix == ".csv":
        return extract_csv_text(path)
    if suffix in {".txt", ".md", ".xml"}:
        return extract_plain_text(path)
    if suffix == ".json":
        return extract_json_text(path)
    if suffix in {".html", ".htm"}:
        return html_to_text(path.read_text(encoding="utf-8", errors="ignore"))
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
        return extract_image_text(path)
    if suffix == ".zip":
        return extract_zip_text(path)
    return ""


def extract_plain_text(path: Path) -> str:
    return clean_text(path.read_text(encoding="utf-8", errors="ignore"))


def extract_json_text(path: Path) -> str:
    with path.open("r", encoding="utf-8", errors="ignore") as f:
        data = json.load(f)
    return clean_text(json.dumps(data, ensure_ascii=False, indent=2))


def extract_csv_text(path: Path) -> str:
    rows: list[str] = []
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as f:
        reader = csv.reader(f)
        for row in reader:
            rows.append("\t".join(str(cell) for cell in row))
    return clean_text("\n".join(rows))


def extract_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            continue
    return clean_text("\n\n".join(pages))


def extract_docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text]
    return clean_text("\n".join(parts))


def extract_xlsx_text(path: Path) -> str:
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if value is None else str(value) for value in row]
            if any(cells):
                parts.append("\t".join(cells))
    return clean_text("\n".join(parts))


def extract_image_text(path: Path) -> str:
    if not pytesseract or not Image:
        return ""
    try:
        image = Image.open(path)
        return clean_text(pytesseract.image_to_string(image))
    except Exception:
        return ""


def extract_zip_text(path: Path) -> str:
    collected: list[str] = []
    with TemporaryDirectory() as temp_dir_name:
        temp_dir = Path(temp_dir_name)
        with zipfile.ZipFile(path, "r") as archive:
            for member in archive.namelist():
                if member.endswith("/"):
                    continue
                target = temp_dir / Path(member).name
                with archive.open(member) as src, target.open("wb") as dst:
                    shutil.copyfileobj(src, dst)
                collected.append(f"=== {member} ===")
                collected.append(extract_text_from_file(target))
    return clean_text("\n\n".join(part for part in collected if part))


def parse_eml_file(path: Path) -> dict[str, Any]:
    with path.open("rb") as f:
        message = BytesParser(policy=policy.default).parse(f)
    return parse_email_message(message, raw_path=path)


def parse_mbox_file(path: Path) -> Iterator[dict[str, Any]]:
    box = mbox(str(path))
    for item in box:
        yield parse_email_message(item, raw_path=path)


def parse_msg_file(path: Path) -> dict[str, Any]:
    if extract_msg is None:
        raise RuntimeError("extract-msg is not installed")
    msg = extract_msg.Message(str(path))
    attachments: list[dict[str, Any]] = []
    for index, attachment in enumerate(msg.attachments):
        name = getattr(attachment, "longFilename", None) or getattr(attachment, "filename", None) or f"attachment_{index}"
        data = getattr(attachment, "data", b"") or b""
        attachments.append(
            {
                "filename": name,
                "content_type": mimetypes.guess_type(name)[0],
                "data": data,
            }
        )
    html_body = None
    try:
        html_body = msg.htmlBody
    except Exception:
        html_body = None
    return {
        "subject": msg.subject or "",
        "sender": msg.sender or "",
        "recipients": msg.to or "",
        "cc": msg.cc or "",
        "bcc": msg.bcc or "",
        "sent_at": str(getattr(msg, "date", "") or ""),
        "folder": "",
        "body_text": clean_text(msg.body or ""),
        "body_html": html_body.decode("utf-8", errors="ignore") if isinstance(html_body, bytes) else (html_body or ""),
        "attachments": attachments,
        "raw_path": str(path),
    }


def parse_email_message(message: Message, raw_path: Path | None = None) -> dict[str, Any]:
    text_parts: list[str] = []
    html_parts: list[str] = []
    attachments: list[dict[str, Any]] = []

    if message.is_multipart():
        for part in message.walk():
            if part.is_multipart():
                continue
            disposition = part.get_content_disposition()
            content_type = part.get_content_type()
            filename = part.get_filename()
            payload = part.get_payload(decode=True) or b""

            if disposition == "attachment" or filename:
                attachments.append(
                    {
                        "filename": filename or "attachment.bin",
                        "content_type": content_type,
                        "data": payload,
                    }
                )
                continue

            if content_type == "text/plain":
                charset = part.get_content_charset() or "utf-8"
                text_parts.append(payload.decode(charset, errors="ignore"))
            elif content_type == "text/html":
                charset = part.get_content_charset() or "utf-8"
                html_parts.append(payload.decode(charset, errors="ignore"))
    else:
        payload = message.get_payload(decode=True) or b""
        content_type = message.get_content_type()
        charset = message.get_content_charset() or "utf-8"
        if content_type == "text/plain":
            text_parts.append(payload.decode(charset, errors="ignore"))
        elif content_type == "text/html":
            html_parts.append(payload.decode(charset, errors="ignore"))

    body_html = "\n\n".join(html_parts)
    body_text = clean_text("\n\n".join(text_parts))
    if not body_text and body_html:
        body_text = html_to_text(body_html)

    return {
        "subject": message.get("subject", ""),
        "sender": message.get("from", ""),
        "recipients": message.get("to", ""),
        "cc": message.get("cc", ""),
        "bcc": message.get("bcc", ""),
        "sent_at": message.get("date", ""),
        "folder": "",
        "body_text": body_text,
        "body_html": body_html,
        "attachments": attachments,
        "raw_path": str(raw_path) if raw_path else None,
    }


def run_readpst(pst_path: Path, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    command = [
        "readpst",
        "-e",
        "-8",
        "-j",
        "0",
        "-q",
        "-w",
        "-o",
        str(output_dir),
        str(pst_path),
    ]
    subprocess.run(command, check=True)
    return output_dir
